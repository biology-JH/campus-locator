#!/usr/bin/env python3
"""Convert the supplied teacher-abbreviation XLSX into a compact JSON lookup.

Requires openpyxl (data-preparation dependency only; not used by the web server).
"""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def title_word(value: str) -> str:
    return "-".join(part[:1].upper() + part[1:].lower() for part in value.split("-"))


def display_name(source_name: str) -> str:
    parts = source_name.strip().split()
    if not parts:
        return source_name.strip()

    surname_count = 0
    for part in parts:
        letters = re.sub(r"[^A-Za-zÀ-ÖØ-öø-ÿ]", "", part)
        if letters and letters == letters.upper() and letters != letters.lower():
            surname_count += 1
        else:
            break

    if surname_count == 0 or surname_count == len(parts):
        return " ".join(title_word(part) for part in parts)

    given = " ".join(title_word(part) for part in parts[surname_count:])
    surname = " ".join(title_word(part) for part in parts[:surname_count])
    return f"{given} {surname}"


def convert(workbook_path: Path) -> dict:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Teachers" not in workbook.sheetnames:
        raise ValueError("Expected a worksheet named 'Teachers'.")

    sheet = workbook["Teachers"]
    header_row = None
    name_column = short_column = None
    for row_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 30), values_only=True), start=1):
        normalized = {str(value).strip().lower(): index for index, value in enumerate(row) if value is not None}
        if "name" in normalized and "short" in normalized:
            header_row = row_index
            name_column = normalized["name"]
            short_column = normalized["short"]
            break

    if header_row is None or name_column is None or short_column is None:
        raise ValueError("Could not find the 'Name' and 'Short' columns in the Teachers worksheet.")

    staff = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if name_column >= len(row) or short_column >= len(row):
            continue
        raw_name, raw_code = row[name_column], row[short_column]
        if raw_name in (None, "") or raw_code in (None, ""):
            continue
        name = str(raw_name).strip()
        code = str(raw_code).strip().upper()
        if not re.fullmatch(r"[A-Z0-9]{1,8}", code):
            continue
        staff.append({"code": code, "sourceName": name, "displayName": display_name(name)})

    by_code: dict[str, list[dict]] = defaultdict(list)
    for entry in staff:
        by_code[entry["code"]].append(
            {"displayName": entry["displayName"], "sourceName": entry["sourceName"]}
        )

    return {
        "sourceFile": workbook_path.name,
        "staffCount": len(staff),
        "uniqueCodeCount": len(by_code),
        "staff": staff,
        "byCode": dict(sorted(by_code.items())),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="Teacher abbreviation XLSX")
    parser.add_argument("output", type=Path, help="Destination JSON lookup")
    args = parser.parse_args()

    data = convert(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(data, ensure_ascii=False, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )
    ambiguous = sum(len(matches) > 1 for matches in data["byCode"].values())
    print(
        f"Extracted {data['staffCount']} staff records, {data['uniqueCodeCount']} codes, "
        f"and {ambiguous} ambiguous codes into {args.output}"
    )


if __name__ == "__main__":
    main()
