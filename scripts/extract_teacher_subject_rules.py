#!/usr/bin/env python3
"""Extract grade/course teacher rules from a school Lessons workbook.

Usage: python3 scripts/extract_teacher_subject_rules.py lessons.xlsx output.json
The output joins workbook teacher names to the normalized abbreviation lookup.
"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def normalized_name(value: str) -> tuple[str, ...]:
    ascii_value = (
        unicodedata.normalize("NFKD", value)
        .encode("ascii", "ignore")
        .decode("ascii")
        .lower()
    )
    return tuple(sorted(re.findall(r"[a-z0-9]+", ascii_value)))


def extract_rules(workbook_path: Path, abbreviations_path: Path) -> dict:
    abbreviation_data = json.loads(abbreviations_path.read_text(encoding="utf-8"))
    names_by_key: dict[tuple[str, ...], list[tuple[str, str]]] = defaultdict(list)
    for code, people in abbreviation_data.get("byCode", {}).items():
        for person in people:
            names_by_key[normalized_name(person["displayName"])].append(
                (code.upper(), person["displayName"])
            )

    worksheet = load_workbook(workbook_path, read_only=True, data_only=True).active
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    indices = {header: index for index, header in enumerate(headers)}
    required = {"Teacher", "Subject"}
    if not required.issubset(indices):
        raise ValueError(f"Lessons sheet must contain columns {sorted(required)}")

    records: dict[str, dict[tuple[str, str], set[str]]] = defaultdict(
        lambda: defaultdict(set)
    )
    matched_rows = 0
    unmatched_rows = 0

    for row in rows:
        teacher_cell = str(row[indices["Teacher"]] or "").strip()
        subject_cell = str(row[indices["Subject"]] or "").strip().upper()
        course_match = re.match(r"^(\d{2})([A-Z]+)", subject_cell)
        if not course_match:
            continue

        grade, course = course_match.groups()
        for suffix in ("HL", "SL"):
            if course.endswith(suffix):
                course = course[: -len(suffix)]
                break

        teacher_matches: list[tuple[str, str]] = []
        for teacher_name in (part.strip() for part in teacher_cell.split(",")):
            if not teacher_name:
                continue
            candidates = names_by_key.get(normalized_name(teacher_name), [])
            if len(candidates) == 1:
                teacher_matches.append(candidates[0])

        if not teacher_matches:
            unmatched_rows += 1
            continue

        matched_rows += 1
        for code, display_name in teacher_matches:
            records[code][(grade, course)].add(display_name)

    by_code = {
        code: [
            {"grade": grade, "subject": subject, "names": sorted(names)}
            for (grade, subject), names in sorted(course_records.items())
        ]
        for code, course_records in sorted(records.items())
    }
    return {
        "sourceFile": workbook_path.name,
        "matchedLessonRows": matched_rows,
        "unmatchedLessonRows": unmatched_rows,
        "byCode": by_code,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="Lessons workbook (.xlsx)")
    parser.add_argument("output", type=Path, help="Destination JSON file")
    parser.add_argument(
        "--abbreviations",
        type=Path,
        default=Path(__file__).resolve().parents[1]
        / "server/data/teacher-abbreviations.json",
        help="Normalized staff abbreviation lookup JSON",
    )
    args = parser.parse_args()

    data = extract_rules(args.workbook, args.abbreviations)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(data, ensure_ascii=False, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )
    print(
        f"Matched {data['matchedLessonRows']} lesson rows; "
        f"{data['unmatchedLessonRows']} rows did not match the abbreviation lookup; "
        f"wrote rules for {len(data['byCode'])} teacher codes to {args.output}"
    )


if __name__ == "__main__":
    main()
